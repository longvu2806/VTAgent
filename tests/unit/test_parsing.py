import docx
import pymupdf
import pytest

from core.parsing import UnsupportedDocumentError, parse


def test_pdf_keeps_page_numbers_and_skips_blank_pages(tmp_path):
    path = tmp_path / "a.pdf"
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "First page")
    pdf.new_page()
    pdf.new_page().insert_text((72, 72), "Third page")
    pdf.save(path)

    pages = parse(path)

    assert [p.page for p in pages] == [1, 3]
    assert "First page" in pages[0].text
    assert "Third page" in pages[1].text


def test_docx_keeps_paragraphs_and_tables_in_order(tmp_path):
    path = tmp_path / "a.docx"
    document = docx.Document()
    document.add_paragraph("Điều 1. Giá trị hợp đồng")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("Hạng mục", "Giá"), ("Thép", "2.500.000")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("Điều 2. Phạt vi phạm")
    document.save(path)

    [page] = parse(path)

    assert page.page is None
    assert "| Hạng mục | Giá |\n| --- | --- |\n| Thép | 2.500.000 |" in page.text
    assert page.text.index("Điều 1") < page.text.index("| Thép") < page.text.index("Điều 2")


def test_pdf_without_text_layer_is_rejected(tmp_path):
    path = tmp_path / "scan.pdf"
    pdf = pymupdf.open()
    pdf.new_page()
    pdf.save(path)

    with pytest.raises(UnsupportedDocumentError, match="no text layer"):
        parse(path)


def test_unsupported_format_is_rejected(tmp_path):
    path = tmp_path / "a.xlsx"
    path.write_bytes(b"x")

    with pytest.raises(UnsupportedDocumentError, match="unsupported format"):
        parse(path)
