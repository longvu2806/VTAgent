"""Ingest + ask end to end, with fake models and in-memory Qdrant."""

from functools import partial
from pathlib import Path

import docx
import pytest
from langchain_core.embeddings import DeterministicFakeEmbedding
from langchain_core.language_models import FakeListChatModel
from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
from langchain_core.messages import AIMessage, AIMessageChunk
from langchain_core.outputs import ChatGenerationChunk
from qdrant_client import QdrantClient

from core.chunking import split_pages_fixed, split_structured
from core.context import assemble
from core.ingestion import Ingestor
from core.pipeline import RAGPipeline
from core.prompting import NOT_FOUND
from core.rerank import Reranker
from core.retrieval import Retriever
from core.sparse import BM25Encoder
from core.stores import DocumentRegistry, VectorStore


class OllamaLikeChatModel(GenericFakeChatModel):
    """Streams like ChatOllama: the text first, then durations and token usage on a last chunk."""

    reply_metadata: dict
    reply_usage: dict

    def _stream(self, messages, stop=None, run_manager=None, **kwargs):
        yield from super()._stream(messages, stop, run_manager, **kwargs)
        yield ChatGenerationChunk(message=AIMessageChunk(
            content="", response_metadata=self.reply_metadata, usage_metadata=self.reply_usage))


class CountingEmbedding(DeterministicFakeEmbedding):
    document_calls: int = 0

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        self.document_calls += 1
        return super().embed_documents(texts)


@pytest.fixture
def client():
    return QdrantClient(":memory:")


@pytest.fixture
def embedder():
    return CountingEmbedding(size=16)


@pytest.fixture
def contract(tmp_path):
    path = tmp_path / "hd.docx"
    document = docx.Document()
    document.add_paragraph("Điều 5. Bên B chịu phạt 0,05% mỗi ngày chậm giao hàng.")
    document.save(path)
    return path


def _rag(embedder, client, llm, collection="test", timeout_s=None, **retriever_options):
    return RAGPipeline(Retriever(embedder, VectorStore(client, collection), top_k=5, **retriever_options), llm,
                       timeout_s=timeout_s)


def _ingestor(tmp_path, tokenizer, embedder, client, collection="test", chunker=None):
    return Ingestor(
        files_dir=tmp_path / "files",
        chunker=chunker or partial(split_pages_fixed, tokenizer=tokenizer, chunk_size=400, overlap=50),
        embedder=embedder,
        store=VectorStore(client, collection),
        registry=DocumentRegistry(tmp_path / "db.sqlite", collection),
    )


def test_ingest_then_ask_maps_citation_to_chunk(tmp_path, word_tokenizer, embedder, client, contract):
    record = _ingestor(tmp_path, word_tokenizer, embedder, client).ingest(contract)

    assert record.n_chunks == 1
    assert Path(record.stored_path).read_bytes() == contract.read_bytes()

    llm = FakeListChatModel(responses=["Phạt 0,05% mỗi ngày [S1]."])
    answer = _rag(embedder, client, llm).ask("Mức phạt chậm giao hàng?")

    assert answer.text == "Phạt 0,05% mỗi ngày [S1]."
    [citation] = answer.citations
    assert citation.chunk.doc_id == record.doc_id
    assert citation.chunk.source == "hd.docx"
    assert "0,05%" in citation.chunk.text


def test_structured_children_reach_the_llm_as_their_article(tmp_path, word_tokenizer, embedder, client):
    path = tmp_path / "luat.docx"
    document = docx.Document()
    for line in ["Điều 8. Các hành vi bị nghiêm cấm", "1. Phá hoại đường, cầu.", "2. Lái xe ô tô có nồng độ cồn."]:
        document.add_paragraph(line)
    document.save(path)
    chunker = partial(split_structured, tokenizer=word_tokenizer, child_min=1, child_max=400, parent_max=1200,
                      fallback_max=300)
    assert _ingestor(tmp_path, word_tokenizer, embedder, client, chunker=chunker).ingest(path).n_chunks == 2

    llm = FakeListChatModel(responses=["Bị nghiêm cấm [S1]."])
    rag = _rag(embedder, client, llm,
               assemble=partial(assemble, budget=2500, count_tokens=lambda t: len(t.split())))
    answer = rag.ask("Lái xe có nồng độ cồn có bị cấm không?")

    assert len(answer.retrieved) == 2 and answer.retrieved[0].chunk.parent_id == answer.retrieved[1].chunk.parent_id
    [source] = answer.sources
    assert "Phá hoại" in source.chunk.text and "nồng độ cồn" in source.chunk.text
    assert source.chunk.section == "Điều 8. Các hành vi bị nghiêm cấm"
    assert answer.citations[0].chunk == source.chunk
    assert "assemble" in answer.timings


def test_reingesting_same_file_skips_embedding(tmp_path, word_tokenizer, embedder, client, contract):
    ingestor = _ingestor(tmp_path, word_tokenizer, embedder, client)
    first = ingestor.ingest(contract)
    second = ingestor.ingest(contract)

    assert second == first
    assert embedder.document_calls == 1


def test_same_file_is_indexed_again_for_another_collection(tmp_path, word_tokenizer, embedder, client, contract):
    _ingestor(tmp_path, word_tokenizer, embedder, client, "r0").ingest(contract)
    _ingestor(tmp_path, word_tokenizer, embedder, client, "r1").ingest(contract)

    assert embedder.document_calls == 2


def _hybrid(tmp_path, tokenizer, embedder, client, contract, llm, score):
    """Ingest with BM25 sparse vectors, then a pipeline with hybrid search and reranking (R4)."""
    bm25 = BM25Encoder(k1=1.2, b=0.75, avg_len=10)
    Ingestor(tmp_path / "files", partial(split_pages_fixed, tokenizer=tokenizer, chunk_size=400, overlap=50),
             embedder, VectorStore(client, "r3", sparse_idf=True), DocumentRegistry(tmp_path / "db.sqlite", "r3"),
             sparse_encoder=bm25).ingest(contract)
    retriever = Retriever(embedder, VectorStore(client, "r3", sparse_idf=True), top_k=5, sparse_encoder=bm25,
                          rerank=Reranker(score, keep=5, threshold=0.5), candidates=40)
    return RAGPipeline(retriever, llm)


def test_hybrid_search_then_rerank_answers_with_the_reranker_scores(tmp_path, word_tokenizer, embedder, client,
                                                                     contract):
    llm = FakeListChatModel(responses=["Phạt 0,05% [S1]."])
    rag = _hybrid(tmp_path, word_tokenizer, embedder, client, contract, llm, lambda q, ps: [0.8] * len(ps))

    answer = rag.ask("Phạt 0,05% mỗi ngày?")

    assert answer.text == "Phạt 0,05% [S1]." and answer.retrieved[0].score == 0.8
    assert {"embed_query", "sparse_query", "search", "rerank", "generate"} <= set(answer.timings)


def test_nothing_above_the_rerank_threshold_answers_not_found_without_calling_llm(tmp_path, word_tokenizer, embedder,
                                                                                   client, contract):
    llm = FakeListChatModel(responses=[])  # raises if invoked
    rag = _hybrid(tmp_path, word_tokenizer, embedder, client, contract, llm, lambda q, ps: [0.1] * len(ps))

    answer = rag.ask("Trọng tài ở đâu?")

    assert (answer.text, answer.sources) == (NOT_FOUND, [])
    assert "rerank" in answer.timings


def test_empty_index_answers_not_found_without_calling_llm(embedder, client):
    llm = FakeListChatModel(responses=[])  # raises if invoked
    answer = _rag(embedder, client, llm, "empty").ask("Bất kỳ?")

    assert answer.text == NOT_FOUND
    assert answer.citations == []
    assert answer.sources == []
    assert set(answer.timings) == {"embed_query", "search"}


class RecordingChatModel(FakeListChatModel):
    """Keeps the messages of every call."""

    calls: list = []

    def _stream(self, messages, stop=None, run_manager=None, **kwargs):
        self.calls.append(messages)
        yield from super()._stream(messages, stop, run_manager, **kwargs)


def test_the_pipeline_sends_its_system_prompt(tmp_path, word_tokenizer, embedder, client, contract):
    _ingestor(tmp_path, word_tokenizer, embedder, client).ingest(contract)
    llm = RecordingChatModel(responses=["Phạt 0,05% [S1]."], calls=[])

    RAGPipeline(Retriever(embedder, VectorStore(client, "test"), top_k=5), llm, system_prompt="Answer briefly.").ask("Phạt?")

    [messages] = llm.calls
    assert (messages[0].type, messages[0].content) == ("system", "Answer briefly.")


def test_ask_reports_stage_timings_and_ollama_stats(tmp_path, word_tokenizer, embedder, client, contract):
    _ingestor(tmp_path, word_tokenizer, embedder, client).ingest(contract)
    llm = OllamaLikeChatModel(
        messages=iter([AIMessage(content="Phạt 0,05% [S1].")]),
        reply_metadata={"load_duration": 1e9, "prompt_eval_duration": 2.5e9, "eval_duration": 4e9},
        reply_usage={"input_tokens": 1700, "output_tokens": 60, "total_tokens": 1760},
    )

    answer = _rag(embedder, client, llm).ask("Mức phạt?")

    assert {"embed_query", "search", "generate"} <= set(answer.timings)
    assert answer.timings["llm_load"] == 1.0
    assert answer.timings["llm_prefill"] == 2.5
    assert answer.timings["llm_decode"] == 4.0
    assert answer.tokens == {"input_tokens": 1700, "output_tokens": 60}
    assert not answer.timed_out


def test_llm_is_stopped_at_the_time_limit(tmp_path, word_tokenizer, embedder, client, contract):
    _ingestor(tmp_path, word_tokenizer, embedder, client).ingest(contract)
    looping = AIMessage(content="Không được [S1]. " + "Tuy nhiên, xét lại [S1]. " * 50)
    llm = GenericFakeChatModel(messages=iter([looping]))

    answer = _rag(embedder, client, llm, timeout_s=0).ask("Mức phạt?")

    assert answer.timed_out
    assert 0 < len(answer.text) < len(looping.content)
    assert answer.citations == [] or answer.citations[0].label == "S1"
